import os
from flask import Flask, render_template, request, jsonify, send_file
from werkzeug.utils import secure_filename
import spacy
from PyPDF2 import PdfReader
from docx import Document
import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import language_tool_python
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.graphics.shapes import Drawing, Rect
from reportlab.graphics import renderPDF
from io import BytesIO
from datetime import datetime
import re
import colorsys

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16MB max file size

# Ensure upload directory exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Load SpaCy model
nlp = spacy.load('en_core_web_sm')

# Initialize language tool
tool = language_tool_python.LanguageTool('en-US')

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in {'pdf', 'docx'}

def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def extract_keywords(text):
    doc = nlp(text)
    keywords = []
    for token in doc:
        if token.pos_ in ['NOUN', 'PROPN'] and not token.is_stop:
            keywords.append(token.text.lower())
    return list(set(keywords))

def calculate_match_percentage(resume_text, job_description):
    vectorizer = TfidfVectorizer()
    try:
        tfidf_matrix = vectorizer.fit_transform([resume_text, job_description])
        similarity = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]
        return round(similarity * 100, 2)
    except:
        return 0

def check_grammar(text):
    matches = tool.check(text)
    suggestions = []
    for match in matches:
        if match.ruleId not in ['UPPERCASE_SENTENCE_START', 'EN_QUOTES']:  # Ignore some common rules
            suggestions.append({
                'context': match.context,
                'message': match.message,
                'replacements': match.replacements[:3]  # Get up to 3 suggested replacements
            })
    return suggestions

def analyze_ats_compatibility(resume_text):
    """Analyze resume for ATS compatibility."""
    ats_issues = []
    recommendations = []
    
    # Check for common ATS issues
    doc = nlp(resume_text)
    
    # Check for headers
    headers = [token.text for token in doc if token.text.upper() == token.text and len(token.text.split()) <= 3]
    required_headers = ['EXPERIENCE', 'EDUCATION', 'SKILLS']
    missing_headers = [h for h in required_headers if h not in headers]
    if missing_headers:
        ats_issues.append({
            'type': 'missing_sections',
            'message': f"Missing standard sections: {', '.join(missing_headers)}",
            'recommendation': "Add standard section headers for better ATS parsing"
        })
    
    # Check for keywords density
    total_words = len([token for token in doc if token.is_alpha])
    if total_words < 200:
        ats_issues.append({
            'type': 'content_length',
            'message': "Resume is too short for optimal ATS scanning",
            'recommendation': "Add more detailed descriptions of your experience and skills"
        })
    
    # Check for formatting issues
    if '\t' in resume_text:
        ats_issues.append({
            'type': 'formatting',
            'message': "Tabs detected in resume",
            'recommendation': "Replace tabs with spaces for better ATS compatibility"
        })
    
    # Check for special characters
    special_chars = re.findall(r'[^\w\s.,;:()\-/]', resume_text)
    if special_chars:
        ats_issues.append({
            'type': 'special_characters',
            'message': "Special characters detected that might affect ATS parsing",
            'recommendation': "Remove or replace special characters with standard alternatives"
        })
    
    return ats_issues

def generate_enhanced_resume(resume_text, job_description, analysis_data, customization_options=None):
    """Generate an enhanced version of the resume based on analysis and customization options."""
    print("Starting enhanced resume generation...")
    try:
        if not resume_text or not isinstance(resume_text, str):
            print("Invalid resume text provided")
            return resume_text

        if customization_options is None:
            customization_options = {
                'add_missing_keywords': True,
                'enhance_experience': True,
                'optimize_summary': True,
                'ats_optimization': True,
                'keyword_density': 'medium',
                'style': 'professional'
            }
        
        print(f"Customization options: {customization_options}")
        
        # Split resume into sections
        sections = {}
        current_section = "DEFAULT"
        sections[current_section] = []
        
        # Split by common section headers
        section_headers = ['EXPERIENCE', 'WORK EXPERIENCE', 'EDUCATION', 'SKILLS', 
                         'TECHNICAL SKILLS', 'SUMMARY', 'PROFESSIONAL SUMMARY', 
                         'PROJECTS', 'CERTIFICATIONS']
        
        lines = resume_text.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Check if line is a section header
            is_header = False
            for header in section_headers:
                if line.upper() == header or line.upper().startswith(header):
                    current_section = header
                    sections[current_section] = []
                    is_header = True
                    break
            
            if not is_header and line:
                if current_section not in sections:
                    sections[current_section] = []
                sections[current_section].append(line)
        
        print(f"Found sections: {list(sections.keys())}")
        
        # Enhance each section
        enhanced_sections = []
        for section, content in sections.items():
            enhanced_content = []
            
            # Enhance Experience section
            if section.upper() in ['EXPERIENCE', 'WORK EXPERIENCE'] and customization_options['enhance_experience']:
                print(f"Enhancing experience section: {section}")
                for line in content:
                    enhanced_line = line
                    if customization_options['add_missing_keywords'] and 'missing_keywords' in analysis_data:
                        for keyword in analysis_data['missing_keywords']:
                            if keyword not in line.lower():
                                if customization_options['style'] == 'technical':
                                    enhanced_line += f" Expertise in {keyword.title()}."
                                elif customization_options['style'] == 'creative':
                                    enhanced_line += f" Demonstrated excellence in {keyword}."
                                else:  # professional
                                    enhanced_line += f" Strong {keyword} skills."
                    enhanced_content.append(enhanced_line)
            
            # Enhance Skills section
            elif section.upper() in ['SKILLS', 'TECHNICAL SKILLS']:
                print(f"Enhancing skills section: {section}")
                skills = []
                for line in content:
                    skills.extend([s.strip() for s in line.split(',')])
                
                if customization_options['add_missing_keywords'] and 'missing_keywords' in analysis_data:
                    missing_skills = [kw for kw in analysis_data['missing_keywords'] 
                                    if kw not in ' '.join(skills).lower()]
                    skills.extend(missing_skills)
                
                # Organize skills based on style
                if customization_options['style'] == 'technical':
                    skills = sorted(skills, key=lambda x: len(x), reverse=True)
                elif customization_options['style'] == 'creative':
                    skills = sorted(skills)
                else:  # professional
                    skills = sorted(skills, key=lambda x: x.lower())
                
                enhanced_content.append(', '.join(skills))
            
            # Enhance Summary section
            elif section.upper() in ['SUMMARY', 'PROFESSIONAL SUMMARY'] and customization_options['optimize_summary']:
                print(f"Enhancing summary section: {section}")
                if content:
                    summary = content[0]
                    if customization_options['add_missing_keywords'] and 'missing_keywords' in analysis_data:
                        for keyword in analysis_data['missing_keywords'][:3]:
                            if keyword not in summary.lower():
                                if customization_options['style'] == 'technical':
                                    summary = f"{summary} Expert in {keyword}."
                                elif customization_options['style'] == 'creative':
                                    summary = f"{summary} Passionate about {keyword}."
                                else:  # professional
                                    summary = f"{summary} Skilled in {keyword}."
                    enhanced_content.append(summary)
            
            # Keep other sections as is
            else:
                enhanced_content.extend(content)
            
            enhanced_sections.append((section, enhanced_content))
        
        # Generate the enhanced resume text
        enhanced_resume = []
        for section, content in enhanced_sections:
            if section != "DEFAULT":  # Don't add DEFAULT as a section header
                enhanced_resume.append(section)
            enhanced_resume.extend(content)
            enhanced_resume.append('')  # Add blank line between sections
        
        # Apply ATS optimization if requested
        if customization_options['ats_optimization']:
            print("Applying ATS optimization")
            enhanced_resume = [line.replace('\t', '    ') for line in enhanced_resume]  # Replace tabs
            enhanced_resume = [re.sub(r'[^\w\s.,;:()\-/]', '', line) for line in enhanced_resume]  # Remove special chars
        
        result = '\n'.join(enhanced_resume)
        print(f"Enhanced resume generated, length: {len(result)}")
        return result
        
    except Exception as e:
        print(f"Error generating enhanced resume: {str(e)}")
        return resume_text  # Return original resume if enhancement fails

def generate_pdf_report(analysis_data, original_resume, enhanced_resume):
    """Generate a comprehensive PDF report with analysis and enhanced resume."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=16,
        spaceAfter=30
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=14,
        spaceAfter=20
    )
    
    normal_style = ParagraphStyle(
        'CustomNormal',
        parent=styles['Normal'],
        fontSize=12,
        spaceAfter=10
    )
    
    # Content
    content = []
    
    # Title
    content.append(Paragraph("Resume Analysis Report", title_style))
    content.append(Paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", normal_style))
    content.append(Spacer(1, 20))
    
    # Match Percentage
    content.append(Paragraph("Match Percentage", heading_style))
    content.append(Paragraph(f"Your resume matches {analysis_data['match_percentage']}% with the job description.", normal_style))
    content.append(Spacer(1, 20))
    
    # Keywords Analysis
    content.append(Paragraph("Keywords Analysis", heading_style))
    content.append(Paragraph("Found Keywords:", normal_style))
    content.append(Paragraph(", ".join(analysis_data['keywords']), normal_style))
    content.append(Spacer(1, 10))
    
    content.append(Paragraph("Missing Keywords:", normal_style))
    content.append(Paragraph(", ".join(analysis_data['missing_keywords']), normal_style))
    content.append(Spacer(1, 20))
    
    # Grammar Suggestions
    if analysis_data['grammar_suggestions']:
        content.append(Paragraph("Grammar Suggestions", heading_style))
        for suggestion in analysis_data['grammar_suggestions']:
            content.append(Paragraph(f"Context: {suggestion['context']}", normal_style))
            content.append(Paragraph(f"Message: {suggestion['message']}", normal_style))
            if suggestion['replacements']:
                content.append(Paragraph(f"Suggested: {', '.join(suggestion['replacements'])}", normal_style))
            content.append(Spacer(1, 10))
    
    # General Suggestions
    content.append(Paragraph("General Suggestions", heading_style))
    for suggestion in analysis_data['suggestions']:
        content.append(Paragraph(suggestion, normal_style))
    content.append(Spacer(1, 20))
    
    # Enhanced Resume
    content.append(Paragraph("Enhanced Resume", heading_style))
    content.append(Paragraph("Below is an enhanced version of your resume with improvements based on the job description:", normal_style))
    content.append(Spacer(1, 10))
    
    # Split enhanced resume into paragraphs
    for paragraph in enhanced_resume.split('\n\n'):
        if paragraph.strip():
            content.append(Paragraph(paragraph, normal_style))
    
    doc.build(content)
    buffer.seek(0)
    return buffer

def analyze_resume(resume_text, job_description):
    # Extract keywords
    resume_keywords = extract_keywords(resume_text)
    job_keywords = extract_keywords(job_description)
    
    # Calculate match percentage
    match_percentage = calculate_match_percentage(resume_text, job_description)
    
    # Find missing keywords
    missing_keywords = [kw for kw in job_keywords if kw not in resume_keywords]
    
    # Check grammar
    grammar_suggestions = check_grammar(resume_text)
    
    # Generate suggestions
    suggestions = []
    if len(missing_keywords) > 0:
        suggestions.append(f"Consider adding these keywords: {', '.join(missing_keywords[:5])}")
    
    if len(resume_text.split()) < 200:
        suggestions.append("Your resume seems too short. Consider adding more details about your experience and skills.")
    
    if len(grammar_suggestions) > 0:
        suggestions.append("Your resume has some grammatical issues. Please review the suggestions below.")
    
    return {
        'match_percentage': match_percentage,
        'keywords': resume_keywords,
        'missing_keywords': missing_keywords,
        'grammar_suggestions': grammar_suggestions,
        'suggestions': suggestions
    }

def generate_pdf_resume(resume_data, design_options):
    """Generate a beautiful PDF resume with design customization."""
    print("Starting PDF generation with data:", resume_data)  # Debug log
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    story = []
    
    # Convert hex colors to RGB
    primary_color = colors.HexColor(design_options.get('primary_color', '#2c3e50'))
    secondary_color = colors.HexColor(design_options.get('secondary_color', '#3498db'))
    
    # Create custom styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        textColor=primary_color,
        fontSize=24,
        spaceAfter=20,
        alignment=1  # Center alignment
    )
    
    subtitle_style = ParagraphStyle(
        'CustomSubtitle',
        parent=styles['Heading2'],
        textColor=secondary_color,
        fontSize=16,
        spaceAfter=20,
        alignment=1  # Center alignment
    )
    
    section_style = ParagraphStyle(
        'CustomSection',
        parent=styles['Heading2'],
        textColor=primary_color,
        fontSize=14,
        spaceAfter=10
    )
    
    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['Normal'],
        fontSize=int(design_options.get('font_size', 12)),
        leading=int(design_options.get('line_spacing', 14)),
        spaceAfter=10
    )
    
    # Add header
    header = f"<b>{resume_data.get('name', '')}</b>"
    if resume_data.get('title'):
        header += f"<br/>{resume_data.get('title')}"
    story.append(Paragraph(header, title_style))
    
    # Add contact information
    contact_info = []
    if resume_data.get('email'):
        contact_info.append(f"Email: {resume_data.get('email')}")
    if resume_data.get('phone'):
        contact_info.append(f"Phone: {resume_data.get('phone')}")
    if resume_data.get('location'):
        contact_info.append(f"Location: {resume_data.get('location')}")
    
    if contact_info:
        story.append(Paragraph("<br/>".join(contact_info), subtitle_style))
    
    # Add horizontal line
    story.append(HRFlowable(width="100%", thickness=1, color=primary_color, spaceBefore=10, spaceAfter=20))
    
    # Add summary section
    if resume_data.get('summary'):
        story.append(Paragraph("Summary", section_style))
        story.append(Paragraph(resume_data['summary'], body_style))
        story.append(Spacer(1, 20))
    
    # Add experience section
    if resume_data.get('experience'):
        story.append(Paragraph("Experience", section_style))
        experience_paragraphs = resume_data['experience'].split('\n')
        for para in experience_paragraphs:
            if para.strip():
                story.append(Paragraph(f"• {para}", body_style))
        story.append(Spacer(1, 20))
    
    # Add education section
    if resume_data.get('education'):
        story.append(Paragraph("Education", section_style))
        education_paragraphs = resume_data['education'].split('\n')
        for para in education_paragraphs:
            if para.strip():
                story.append(Paragraph(para, body_style))
        story.append(Spacer(1, 20))
    
    # Add skills section
    if resume_data.get('skills'):
        story.append(Paragraph("Skills", section_style))
        skills = [skill.strip() for skill in resume_data['skills'].split(',')]
        
        if design_options.get('include_skills_chart'):
            # Create skills table
            data = [['Skill', 'Proficiency']]
            for skill in skills:
                data.append([skill, '■■■■■'])
            
            t = Table(data, colWidths=[2*inch, 3*inch])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), secondary_color),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 12),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 10),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]))
            story.append(t)
        else:
            # Create skills list
            skills_text = ", ".join(skills)
            story.append(Paragraph(skills_text, body_style))
    
    # Build PDF
    doc.build(story)
    buffer.seek(0)
    print("PDF generation completed successfully")  # Debug log
    return buffer

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/analyze', methods=['POST'])
def analyze():
    try:
        print("Received analyze request")
        
        if 'resume' not in request.files:
            print("No file in request")
            return jsonify({'error': 'No file uploaded'}), 400
        
        file = request.files['resume']
        job_description = request.form.get('job_description', '')
        customization_options = {
            'add_missing_keywords': request.form.get('add_missing_keywords', 'true').lower() == 'true',
            'enhance_experience': request.form.get('enhance_experience', 'true').lower() == 'true',
            'optimize_summary': request.form.get('optimize_summary', 'true').lower() == 'true',
            'ats_optimization': request.form.get('ats_optimization', 'true').lower() == 'true',
            'keyword_density': request.form.get('keyword_density', 'medium'),
            'style': request.form.get('style', 'professional')
        }
        
        print(f"File received: {file.filename}")
        print(f"Job description length: {len(job_description)}")
        print(f"Customization options: {customization_options}")
        
        if file.filename == '':
            print("Empty filename")
            return jsonify({'error': 'No file selected'}), 400
        
        if not allowed_file(file.filename):
            print(f"Invalid file type: {file.filename}")
            return jsonify({'error': 'Invalid file type. Please upload PDF or DOCX'}), 400
        
        # Check file size (16MB limit)
        file.seek(0, 2)  # Seek to end of file
        file_size = file.tell()
        file.seek(0)  # Reset file pointer
        MAX_FILE_SIZE = 16 * 1024 * 1024  # 16MB in bytes
        
        if file_size > MAX_FILE_SIZE:
            print(f"File too large: {file_size} bytes")
            return jsonify({'error': 'File size exceeds 16MB limit. Please choose a smaller file.'}), 400
        
        if not job_description:
            print("No job description provided")
            return jsonify({'error': 'Job description is required'}), 400
        
        filename = secure_filename(file.filename)
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(file_path)
        
        print(f"File saved to: {file_path}")
        
        try:
            if filename.endswith('.pdf'):
                print("Extracting text from PDF")
                resume_text = extract_text_from_pdf(file_path)
            else:
                print("Extracting text from DOCX")
                resume_text = extract_text_from_docx(file_path)
            
            print(f"Extracted text length: {len(resume_text)}")
            
            if len(resume_text.strip()) == 0:
                print("No text extracted from file")
                return jsonify({'error': 'Could not extract any text from the file. Please ensure the file is not password protected and contains readable text.'}), 400
            
            print("Starting analysis...")
            analysis = analyze_resume(resume_text, job_description)
            print("Analysis completed successfully")
            
            # Analyze ATS compatibility
            ats_analysis = analyze_ats_compatibility(resume_text)
            analysis['ats_issues'] = ats_analysis
            
            # Generate enhanced resume
            print("Generating enhanced resume...")
            enhanced_resume = generate_enhanced_resume(resume_text, job_description, analysis, customization_options)
            print(f"Enhanced resume length: {len(enhanced_resume)}")
            
            # Add enhanced resume to analysis data
            analysis['enhanced_resume'] = enhanced_resume
            analysis['original_resume'] = resume_text
            analysis['customization_options'] = customization_options
            
            # Clean up the uploaded file
            os.remove(file_path)
            print("Temporary file removed")
            
            return jsonify(analysis)
            
        except Exception as e:
            print(f"Error during analysis: {str(e)}")
            # Clean up the file in case of error
            if os.path.exists(file_path):
                os.remove(file_path)
            return jsonify({'error': str(e)}), 500
            
    except Exception as e:
        print(f"Unexpected error: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/download_report', methods=['POST'])
def download_report():
    try:
        analysis_data = request.json
        original_resume = analysis_data.get('original_resume', '')
        enhanced_resume = analysis_data.get('enhanced_resume', '')
        
        pdf_buffer = generate_pdf_report(analysis_data, original_resume, enhanced_resume)
        
        return send_file(
            pdf_buffer,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='resume_analysis_report.pdf'
        )
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/download_enhanced_resume', methods=['POST'])
def download_enhanced_resume():
    try:
        data = request.json
        print("Received resume data:", data)  # Debug log
        
        resume_data = data.get('resume_data', {})
        design_options = data.get('design_options', {})
        
        # Ensure all required fields are present
        required_fields = ['name', 'title', 'summary', 'experience', 'education', 'skills']
        for field in required_fields:
            if not resume_data.get(field):
                print(f"Missing required field: {field}")  # Debug log
                resume_data[field] = f"Sample {field.capitalize()}"  # Provide default value
        
        print("Processed resume data:", resume_data)  # Debug log
        
        pdf_buffer = generate_pdf_resume(resume_data, design_options)
        
        return send_file(
            pdf_buffer,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='enhanced_resume.pdf'
        )
    except Exception as e:
        print("Error generating PDF:", str(e))  # Debug log
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True) 